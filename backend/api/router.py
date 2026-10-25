import os
import shutil
import io
import httpx
from uuid import uuid4

from fastapi import APIRouter, UploadFile, File, Depends, Query, HTTPException
from pydantic import BaseModel
from sqlalchemy.orm import Session
from bs4 import BeautifulSoup
from pypdf import PdfReader
import docx

from shared.db import get_db
from modules.ingest.service.splitter import split_markdown
from modules.ingest.repository.document_repo import save_document
from modules.ingest.domain.models import Document, Chunk
from modules.retrieve.service.search import search
from modules.agent.service.agent_service import ask_question

# 引入所有 manage 模块的 Service
from modules.manage.service.folder_service import (
    get_all_folders, create_folder, add_doc_to_folder, get_documents_in_folder, 
    get_all_documents, get_uncategorized_documents, get_public_documents, 
    get_user_public_docs, toggle_doc_visibility,
    get_public_folders, toggle_favorite_doc, toggle_favorite_folder, get_my_favorites
)

api_router = APIRouter()

# ==================== 1. 知识收录 (ingest) ====================
@api_router.post("/documents/upload", tags=["ingest"])
async def upload_document(
    file: UploadFile = File(...), 
    user_id: str = Query("u_zhangsan"),
    is_public: bool = Query(False),
    db: Session = Depends(get_db)
):
    # 1. 物理保存文件
    upload_dir = "uploads"
    os.makedirs(upload_dir, exist_ok=True) 
    unique_filename = f"{uuid4()}_{file.filename}"
    file_path = os.path.join(upload_dir, unique_filename)
    
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    # 2. 解析文件内容
    filename = file.filename.lower()
    content = ""
    source_type = "file"

    if filename.endswith(".pdf"):
        reader = PdfReader(file_path)
        for page in reader.pages: content += page.extract_text() + "\n"
    elif filename.endswith(".docx"):
        doc = docx.Document(file_path)
        for para in doc.paragraphs: content += para.text + "\n"
    elif filename.endswith((".md", ".txt", ".py", ".java", ".cpp", ".c", ".js", ".ts", ".html", ".css", ".json", ".sql", ".go", ".rs")):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f: content = f.read()
    elif filename.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp")):
        source_type = "image"
        content = f"【图片文件】{file.filename}"
    elif filename.endswith((".mp4", ".webm", ".ogg")):
        source_type = "video"
        content = f"【视频文件】{file.filename}"
    else:
        content = f"【未知文件】{file.filename}"
        
    if not content.strip():
        content = f"【空文件或无法解析：{file.filename}】"
        
    # 3. 存入数据库
    doc_id = str(uuid4())
    doc = Document(
        id=doc_id, 
        title=file.filename, 
        source_type=source_type, 
        source_url=f"/uploads/{unique_filename}", # 保存访问路径
        is_public=is_public, 
        owner_id=user_id
    )
    
    chunks_text = split_markdown(content)
    chunks = [Chunk(id=str(uuid4()), document_id=doc_id, content=text, embedding=None, chunk_index=i) for i, text in enumerate(chunks_text)]
    
    save_document(db, doc, chunks)
    return {"id": doc_id, "title": doc.title, "chunks_count": len(chunks)}

class URLCreate(BaseModel):
    url: str
    is_public: bool = False

@api_router.post("/documents/url", tags=["ingest"])
async def ingest_url(req: URLCreate, user_id: str = Query("u_zhangsan"), db: Session = Depends(get_db)):
    try:
        async with httpx.AsyncClient() as client:
            resp = await client.get(req.url, timeout=10.0)
            resp.raise_for_status()
            soup = BeautifulSoup(resp.text, "html.parser")
            content = soup.get_text(separator="\n", strip=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"抓取失败: {str(e)}")
        
    doc_id = str(uuid4())
    doc = Document(id=doc_id, title=req.url, source_type="url", is_public=req.is_public, owner_id=user_id)
    chunks_text = split_markdown(content)
    chunks = [Chunk(id=str(uuid4()), document_id=doc_id, content=text, embedding=None, chunk_index=i) for i, text in enumerate(chunks_text)]
    
    save_document(db, doc, chunks)
    return {"id": doc_id, "title": doc.title, "chunks_count": len(chunks)}

# ==================== 2. 知识检索 (retrieve) ====================
@api_router.get("/search", tags=["retrieve"])
async def search_docs(q: str = Query(...), top_k: int = 5, db: Session = Depends(get_db)):
    return search(db, q, top_k=top_k)

# ==================== 3. 问答生成 (agent) ====================
class AskRequest(BaseModel):
    question: str
    top_k: int = 3

@api_router.post("/agent/ask", tags=["agent"])
async def agent_chat(req: AskRequest, db: Session = Depends(get_db)):
    return await ask_question(db, req.question, req.top_k)

# ==================== 4. 管理模块 (manage) ====================
class FolderCreate(BaseModel):
    name: str

# --- 文件夹操作 ---
@api_router.get("/folders", tags=["manage"])
async def list_folders(user_id: str = Query("u_zhangsan"), db: Session = Depends(get_db)):
    return get_all_folders(db, user_id)

@api_router.post("/folders", tags=["manage"])
async def add_folder(req: FolderCreate, user_id: str = Query("u_zhangsan"), db: Session = Depends(get_db)):
    return create_folder(db, req.name, user_id)

@api_router.get("/folders/{folder_id}/documents", tags=["manage"])
async def list_folder_docs(folder_id: str, db: Session = Depends(get_db)):
    return get_documents_in_folder(db, folder_id)

@api_router.post("/folders/{folder_id}/documents/{doc_id}", tags=["manage"])
async def move_doc(folder_id: str, doc_id: str, db: Session = Depends(get_db)):
    return add_doc_to_folder(db, doc_id, folder_id)

# --- 文档列表 ---
@api_router.get("/documents/all", tags=["manage"])
async def list_all_docs(db: Session = Depends(get_db)):
    return get_all_documents(db)

@api_router.get("/documents/uncategorized", tags=["manage"])
async def list_uncategorized(db: Session = Depends(get_db)):
    return get_uncategorized_documents(db)

# --- 预览文档内容 ---
@api_router.get("/documents/{doc_id}/content", tags=["manage"])
async def get_document_content(doc_id: str, db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="文件不存在")
    full_content = "\n\n".join([c.content for c in doc.chunks])
    return {"title": doc.title, "content": full_content}

# --- 切换公开/私有 ---
@api_router.post("/documents/{doc_id}/toggle-visibility", tags=["manage"])
async def toggle_visibility(doc_id: str, user_id: str = Query("u_zhangsan"), db: Session = Depends(get_db)):
    return toggle_doc_visibility(db, doc_id, user_id)

# --- 公共广场 ---
@api_router.get("/square", tags=["manage"])
async def get_square(keyword: str = Query(""), db: Session = Depends(get_db)):
    return get_public_documents(db, keyword)

@api_router.get("/square/folders", tags=["manage"])
async def square_folders(db: Session = Depends(get_db)):
    return get_public_folders(db)

@api_router.get("/users/{user_id}/docs", tags=["manage"])
async def user_profile(user_id: str, db: Session = Depends(get_db)):
    return get_user_public_docs(db, user_id)

# --- 收藏功能 ---
@api_router.post("/favorites/doc/{doc_id}", tags=["manage"])
async def fav_doc(doc_id: str, user_id: str = Query("u_zhangsan"), db: Session = Depends(get_db)):
    return toggle_favorite_doc(db, user_id, doc_id)

@api_router.post("/favorites/folder/{folder_id}", tags=["manage"])
async def fav_folder(folder_id: str, user_id: str = Query("u_zhangsan"), db: Session = Depends(get_db)):
    return toggle_favorite_folder(db, user_id, folder_id)

@api_router.get("/favorites", tags=["manage"])
async def my_favorites(user_id: str = Query("u_zhangsan"), db: Session = Depends(get_db)):
    return get_my_favorites(db, user_id)

from modules.manage.service.folder_service import (
    # ... 之前的引入保留 ...
    remove_doc_from_folder, delete_document # 👈 新增
)

@api_router.delete("/folders/{folder_id}/documents/{doc_id}", tags=["manage"])
async def remove_from_folder(folder_id: str, doc_id: str, db: Session = Depends(get_db)):
    """从收藏夹移出文档"""
    return remove_doc_from_folder(db, doc_id, folder_id)

@api_router.delete("/documents/{doc_id}", tags=["manage"])
async def delete_doc(doc_id: str, db: Session = Depends(get_db)):
    """彻底删除文档"""
    return delete_document(db, doc_id)